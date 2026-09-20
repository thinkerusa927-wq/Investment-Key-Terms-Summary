#!/usr/bin/env python3
"""Build a Key Terms Summary DOCX from structured JSON and the bundled template."""

from __future__ import annotations

import argparse
import copy
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def clear_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_text(paragraph, text: str, *, bold: bool = False) -> None:
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold


def get_num_id(paragraph) -> int:
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        raise ValueError("Template model paragraph has no numbering definition")
    return int(p_pr.numPr.numId.val)


def ensure_num_pr(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def add_numbering_instance(doc: Document, source_num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    max_num_id = max(int(n.get(qn("w:numId"))) for n in nums)
    source = next(
        (n for n in nums if int(n.get(qn("w:numId"))) == source_num_id), None
    )
    if source is None:
        raise ValueError(f"Template numbering id {source_num_id} was not found")
    abstract = source.find(qn("w:abstractNumId"))
    if abstract is None:
        raise ValueError("Template numbering instance has no abstract numbering reference")

    new_num_id = max_num_id + 1
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract.get(qn("w:val")))
    new_num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    new_num.append(override)
    numbering.append(new_num)
    return new_num_id


def resize_paragraphs(cell, count: int) -> None:
    if count < 1:
        raise ValueError("Each row must contain at least one item")
    while len(cell.paragraphs) < count:
        cell._tc.append(copy.deepcopy(cell.paragraphs[-1]._p))
    while len(cell.paragraphs) > count:
        cell._tc.remove(cell.paragraphs[-1]._p)


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def load_payload(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    for key in ("title", "date", "rows"):
        if key not in payload:
            raise ValueError(f"Missing required JSON field: {key}")
    if not isinstance(payload["rows"], list) or not payload["rows"]:
        raise ValueError("rows must be a non-empty list")
    for row in payload["rows"]:
        if not row.get("label") or not isinstance(row.get("items"), list) or not row["items"]:
            raise ValueError("Each row requires a label and at least one item")
        for item in row["items"]:
            if "lead" not in item or "body" not in item:
                raise ValueError("Each item requires lead and body fields")
    return payload


def build(template: Path, payload_path: Path, output: Path) -> None:
    payload = load_payload(payload_path)
    if template.resolve() == output.resolve():
        raise ValueError("Output must be different from the template")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)

    doc = Document(output)
    if len(doc.paragraphs) < 2 or not doc.tables or len(doc.tables[0].rows) < 2:
        raise ValueError("Template is missing the expected title/date/table structure")

    set_text(doc.paragraphs[0], str(payload["title"]), bold=True)
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_text(doc.paragraphs[1], str(payload["date"]))
    doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = str(payload["title"])

    table = doc.tables[0]
    model_row = table.rows[1]
    model_tr = copy.deepcopy(model_row._tr)
    row_num_id = get_num_id(model_row.cells[0].paragraphs[0])
    content_source_num_id = get_num_id(model_row.cells[2].paragraphs[0])

    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    mark_repeat_header(table.rows[0])

    for row_data in payload["rows"]:
        table._tbl.append(copy.deepcopy(model_tr))
        row = table.rows[-1]
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        resize_paragraphs(row.cells[0], 1)
        clear_runs(row.cells[0].paragraphs[0])
        ensure_num_pr(row.cells[0].paragraphs[0], row_num_id)

        resize_paragraphs(row.cells[1], 1)
        set_text(row.cells[1].paragraphs[0], str(row_data["label"]), bold=True)
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

        items = row_data["items"]
        resize_paragraphs(row.cells[2], len(items))
        content_num_id = add_numbering_instance(doc, content_source_num_id)
        for paragraph, item in zip(row.cells[2].paragraphs, items):
            clear_runs(paragraph)
            ensure_num_pr(paragraph, content_num_id)
            lead = paragraph.add_run(str(item["lead"]))
            lead.underline = True
            paragraph.add_run(str(item["body"]))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.save(output)
    print(output)


def parse_args() -> argparse.Namespace:
    skill_dir = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--template",
        type=Path,
        default=skill_dir / "assets" / "Key Terms Summary模板.docx",
    )
    parser.add_argument("--input-json", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build(args.template, args.input_json, args.output)


if __name__ == "__main__":
    main()
